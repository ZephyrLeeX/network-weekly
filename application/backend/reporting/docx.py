"""DOCX renderer (W03-T008, SYSTEM_SPEC.md §5/§6).

Renders one :class:`WeeklyReportData` into the fixed 8-section weekly
report with python-docx. Layout is table-first; every value without data
is rendered as ``数据缺失`` — never 0, never interpolated (§6.2). The
renderer only reads the report data structure, which contains no secrets
by construction (§22.2).

File writing follows §5 exactly:

    temporary file in the target directory -> python-docx save ->
    successful close/validation (reopen + 8-heading check) -> atomic replace

so a failed generation can never leave a half-written file behind, and an
existing report for the week stays intact until the new one is complete.

Generation and installation are two separate steps
(`render_report_candidate` + `install_report`; `render_report_docx` is
their one-shot composition): a regenerate first produces a fully
validated *candidate* DOCX under a deterministic side name and only the
explicit install atomically moves it onto the current report path. The
report job commits its database success between the two steps, so any
database failure leaves the previous DOCX bytes untouched (§4.4) and an
interrupted install can be completed from the surviving candidate.
"""

import logging
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

from backend.reporting.period import ReportPeriod, report_file_name
from backend.reporting.service import (
    WeeklyReportData,
)

logger = logging.getLogger(__name__)

MISSING = "数据缺失"
NEVER_RECOVERED = "未恢复"

# §6: the report contains exactly these 8 sections, in this order.
SECTION_TITLES: tuple[str, ...] = (
    "周报基本信息与总体摘要",
    "设备运行状态与掉线/恢复情况",
    "IRF 堆叠成员状态",
    "CPU / 内存统计",
    "重点接口状态与 Down/恢复情况",
    "接口利用率 Top 10 与高利用率异常",
    "CRC / Error / Drop 与 Monitoring Coverage",
    "本周处理问题",
)


class DocxRenderError(RuntimeError):
    """Raised when the generated document fails validation (§5)."""


# Deterministic side name of the not-yet-installed report (§4.4): same
# directory and filesystem as the current DOCX, so installing it stays a
# same-filesystem atomic rename. It never matches the §5 current-report
# name, so it is not downloadable/report-visible state.
CANDIDATE_SUFFIX = ".candidate"


def candidate_file_name(period: ReportPeriod) -> str:
    """The deterministic file name of one week's pending candidate (§4.4)."""

    return f"{report_file_name(period)}{CANDIDATE_SUFFIX}"


# Candidate names are the §5 current-report name plus CANDIDATE_SUFFIX;
# the week code inside follows ReportPeriod.week_code (`YYYY-Www`).
_CANDIDATE_NAME_PATTERN = re.compile(
    r"^network-weekly-report-(\d{4}-W\d{2})\.docx" + re.escape(CANDIDATE_SUFFIX) + r"$"
)


def candidate_week_code(file_name: str) -> str | None:
    """The week code of a candidate file name, None when it is not ours."""

    match = _CANDIDATE_NAME_PATTERN.match(file_name)
    return match.group(1) if match else None


@dataclass(frozen=True)
class RenderedReport:
    """The artifact rendered for one week (one current DOCX per week, §4.4).

    `candidate` is where the validated bytes sit right now; `path` is the
    final current-report location that `weekly_reports.file_path` records
    and that `install_report` atomically moves the candidate onto.
    """

    path: Path
    week_code: str
    candidate: Path


def _fmt_dt(moment: datetime | None, timezone: ZoneInfo) -> str:
    if moment is None:
        return MISSING
    return moment.astimezone(timezone).strftime("%Y-%m-%d %H:%M")


def _fmt_duration(seconds: float | None) -> str:
    if seconds is None:
        return NEVER_RECOVERED
    total_minutes = int(seconds // 60)
    days, rest = divmod(total_minutes, 24 * 60)
    hours, minutes = divmod(rest, 60)
    parts: list[str] = []
    if days:
        parts.append(f"{days}天")
    if hours:
        parts.append(f"{hours}小时")
    if minutes or not parts:
        parts.append(f"{minutes}分钟")
    return "".join(parts)


def _fmt_percent(value: float | None) -> str:
    return MISSING if value is None else f"{value:.2f}%"


def _fmt_number(value: float | int | None, suffix: str = "") -> str:
    if value is None:
        return MISSING
    return f"{value:g}{suffix}"


def _fill_row(row: _Row, values: list[str]) -> None:
    for cell, text in zip(row.cells, values, strict=True):
        cell.text = text


def _add_table(document: DocumentObject, headers: list[str], rows: list[list[str]]) -> Table:
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    _fill_row(table.rows[0], headers)
    for index, values in enumerate(rows, start=1):
        _fill_row(table.rows[index], values)
    return table


def _configure_base_style(document: DocumentObject) -> None:
    """Make the default style render Chinese text predictably."""

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "SimSun")


def _section_heading(document: DocumentObject, title: str) -> None:
    document.add_heading(title, level=1)


def _paragraph(document: DocumentObject, text: str) -> Paragraph:
    return document.add_paragraph(text)


def _basic_info(document: DocumentObject, data: WeeklyReportData) -> None:
    """Section 1: 周报基本信息与总体摘要 (§6.1)."""

    coverage = data.coverage
    if coverage.devices:
        coverage_summary = (
            f"总体 {_fmt_percent(coverage.coverage_percent)}"
            f"（设备 {len(coverage.devices)} 台，SUCCESS {coverage.success}，"
            f"PARTIAL {coverage.partial}，FAILED {coverage.failed}）"
        )
    else:
        # Empty deployment: no coverage evidence exists — 数据缺失, never a
        # percentage computed over zero devices, never a healthy look (§6.2).
        coverage_summary = "未配置设备/数据缺失（设备 0 台，无 Monitoring Coverage 数据）"
    if coverage.below_target:
        coverage_summary += "；数据完整性不足"

    rows = [
        ["报告周编号", data.period.week_code],
        ["统计开始时间", _fmt_dt(data.period.start, data.timezone)],
        ["统计结束时间", _fmt_dt(data.period.end, data.timezone)],
        ["报告生成时间", _fmt_dt(data.generated_at, data.timezone)],
        ["当前总体状态", data.overall_status],
        ["Monitoring Coverage 摘要", coverage_summary],
    ]
    _add_table(document, ["项目", "内容"], rows)
    document.add_paragraph()
    _paragraph(document, data.summary_text)


def _device_incidents(document: DocumentObject, data: WeeklyReportData) -> None:
    """Section 2: 设备运行状态与掉线/恢复情况 (§9.4)."""

    if not data.device_incidents:
        _paragraph(document, "本周无设备掉线。")
        return
    rows = []
    for summary in data.device_incidents:
        for episode in summary.episodes:
            state = "仍处于 Down" if episode.ongoing_at_period_end else "已恢复"
            rows.append(
                [
                    summary.device_name,
                    _fmt_dt(episode.started_at, data.timezone),
                    _fmt_dt(episode.recovered_at, data.timezone),
                    _fmt_duration(episode.duration_seconds),
                    state,
                ]
            )
    _add_table(
        document,
        ["设备", "开始时间", "恢复时间", "持续时间", "状态"],
        rows,
    )


def _irf_section(document: DocumentObject, data: WeeklyReportData) -> None:
    """Section 3: IRF 堆叠成员状态 (§17)."""

    if not data.irf_summaries:
        _paragraph(document, "本周无 IRF 设备（未配置多成员 IRF）。")
        return
    for summary in data.irf_summaries:
        _paragraph(
            document,
            f"{summary.device_name}：期望成员数 {summary.expected_member_count}，"
            + (
                f"当前成员数 {summary.observed_member_count}"
                f"（{_fmt_dt(summary.latest_observation_at, data.timezone)}）"
                if not summary.data_missing
                else "本周无成功观察，成员状态数据缺失"
            ),
        )
        rows = []
        for member in summary.members:
            if member.observations_count == 0:
                rows.append([str(member.member_id), MISSING, MISSING, MISSING])
                continue
            if member.missing_windows:
                windows_text = "；".join(
                    f"{_fmt_dt(window.started_at, data.timezone)} 至 "
                    + (
                        _fmt_dt(window.reappeared_at, data.timezone)
                        if window.reappeared_at is not None
                        else "期末仍未恢复"
                    )
                    for window in member.missing_windows
                )
            else:
                windows_text = "无缺失"
            role = member.latest_role if member.latest_role is not None else MISSING
            rows.append(
                [
                    str(member.member_id),
                    "在位" if member.latest_observed else "缺失",
                    windows_text,
                    role,
                ]
            )
        _add_table(
            document,
            ["成员编号", "期末状态", "缺失时间段", "角色"],
            rows,
        )
        if summary.role_changes:
            change_rows = [
                [
                    str(change.member_id),
                    _fmt_dt(change.observed_at, data.timezone),
                    change.previous_role,
                    change.new_role,
                ]
                for change in summary.role_changes
            ]
            _paragraph(document, "角色变化：")
            _add_table(
                document,
                ["成员编号", "观察时间", "原角色", "新角色"],
                change_rows,
            )
        document.add_paragraph()


def _cpu_memory(document: DocumentObject, data: WeeklyReportData) -> None:
    """Section 4: CPU / 内存统计 (§14)."""

    rows = []
    for device in data.device_resources:
        cpu = device.cpu
        memory = device.memory
        rows.append(
            [
                device.device_name,
                _fmt_number(cpu.average if cpu else None, "%"),
                _fmt_number(cpu.maximum if cpu else None, "%"),
                _fmt_number(cpu.p95 if cpu else None, "%"),
                _fmt_number(memory.average if memory else None, "%"),
                _fmt_number(memory.maximum if memory else None, "%"),
                _fmt_number(memory.p95 if memory else None, "%"),
                str(len(device.cpu_sustained_high)),
                str(len(device.memory_sustained_high)),
            ]
        )
    _add_table(
        document,
        [
            "设备",
            "CPU 平均",
            "CPU 最大",
            "CPU P95",
            "内存 平均",
            "内存 最大",
            "内存 P95",
            "CPU 高负载区间",
            "内存 高负载区间",
        ],
        rows,
    )
    interval_rows = []
    for device in data.device_resources:
        for interval in device.cpu_sustained_high:
            interval_rows.append(
                [
                    device.device_name,
                    "CPU",
                    _fmt_dt(interval.start, data.timezone),
                    _fmt_dt(interval.end, data.timezone),
                    _fmt_duration(interval.duration_seconds),
                ]
            )
        for interval in device.memory_sustained_high:
            interval_rows.append(
                [
                    device.device_name,
                    "内存",
                    _fmt_dt(interval.start, data.timezone),
                    _fmt_dt(interval.end, data.timezone),
                    _fmt_duration(interval.duration_seconds),
                ]
            )
    if interval_rows:
        _paragraph(document, "持续高负载区间：")
        _add_table(
            document,
            ["设备", "指标", "开始时间", "结束时间", "持续时间"],
            interval_rows,
        )


def _interface_incidents(document: DocumentObject, data: WeeklyReportData) -> None:
    """Section 5: 重点接口状态与 Down/恢复情况 (§13)."""

    if not data.interface_incidents:
        _paragraph(document, "本周无重点接口中断。")
        return
    rows = []
    for summary in data.interface_incidents:
        for episode in summary.episodes:
            state = "仍处于 Down" if episode.ongoing_at_period_end else "已恢复"
            rows.append(
                [
                    summary.device_name,
                    summary.display_name,
                    summary.description or "",
                    _fmt_dt(episode.started_at, data.timezone),
                    _fmt_dt(episode.recovered_at, data.timezone),
                    _fmt_duration(episode.duration_seconds),
                    state,
                ]
            )
    _add_table(
        document,
        ["设备", "接口", "描述", "开始时间", "恢复时间", "持续时间", "状态"],
        rows,
    )


def _utilization(document: DocumentObject, data: WeeklyReportData) -> None:
    """Section 6: 接口利用率 Top 10 与高利用率异常 (§15.4/§15.3)."""

    if not data.interface_top10:
        _paragraph(document, "本周无有效接口利用率数据。")
    else:
        rows = [
            [
                str(index + 1),
                entry.device_name,
                entry.display_name,
                entry.description or "",
                str(entry.sample_count),
                _fmt_number(entry.p95, "%"),
                _fmt_number(entry.maximum, "%"),
            ]
            for index, entry in enumerate(data.interface_top10)
        ]
        _add_table(
            document,
            ["排名", "设备", "接口", "描述", "有效样本", "利用率 P95", "最大利用率"],
            rows,
        )
    document.add_paragraph()
    if not data.high_utilization:
        _paragraph(document, "本周无重点接口持续高利用率。")
        return
    rows = []
    for entry in data.high_utilization:
        for interval in entry.intervals:
            rows.append(
                [
                    entry.device_name,
                    entry.display_name,
                    _fmt_dt(interval.start, data.timezone),
                    _fmt_dt(interval.end, data.timezone),
                    _fmt_duration(interval.duration_seconds),
                ]
            )
    _add_table(
        document,
        ["设备", "接口", "开始时间", "结束时间", "持续时间"],
        rows,
    )


def _counters_coverage(document: DocumentObject, data: WeeklyReportData) -> None:
    """Section 7: CRC / Error / Drop 与 Monitoring Coverage (§16/§18)."""

    if not data.counter_top10:
        _paragraph(document, "本周无错误计数器数据。")
    else:
        rows = [
            [
                entry.device_name,
                entry.display_name,
                _fmt_number(entry.crc_delta),
                _fmt_number(entry.in_errors_delta),
                _fmt_number(entry.out_errors_delta),
                _fmt_number(entry.in_discards_delta),
                _fmt_number(entry.out_discards_delta),
                _fmt_number(entry.total_delta),
            ]
            for entry in data.counter_top10
        ]
        _add_table(
            document,
            ["设备", "接口", "CRC 增量", "入错误", "出错误", "入丢弃", "出丢弃", "合计"],
            rows,
        )
        _paragraph(document, "以上为观察信息，不改变总体状态。")
    document.add_paragraph()

    coverage = data.coverage
    rows = [
        [
            device.device_name,
            str(device.expected),
            str(device.success),
            str(device.partial),
            str(device.failed),
            _fmt_percent(device.coverage_percent),
            "数据完整性不足" if device.below_target else "",
        ]
        for device in coverage.devices
    ]
    overall = [
        "总体",
        str(coverage.expected),
        str(coverage.success),
        str(coverage.partial),
        str(coverage.failed),
        _fmt_percent(coverage.coverage_percent),
        "数据完整性不足" if coverage.below_target else "",
    ]
    _add_table(
        document,
        ["设备", "期望周期", "成功", "部分成功", "失败", "覆盖率", "备注"],
        [*rows, overall],
    )


def _open_questions(document: DocumentObject) -> None:
    """Section 8: 本周处理问题 — clear blank editable area (§6)."""

    _paragraph(document, "（请下载本报告后在此手工填写本周处理的问题。）")
    area = document.add_table(rows=1, cols=1)
    area.style = "Table Grid"
    cell = area.rows[0].cells[0]
    cell.text = ""
    for _ in range(6):
        cell.add_paragraph("")


def _build_document(data: WeeklyReportData) -> DocumentObject:
    document = Document()
    _configure_base_style(document)

    document.add_heading(
        f"网络运维周报 {data.period.week_code}", level=0
    )
    document.add_paragraph(
        f"统计周期：{_fmt_dt(data.period.start, data.timezone)} 至 "
        f"{_fmt_dt(data.period.end, data.timezone)}（Asia/Shanghai）"
    )

    for index, title in enumerate(SECTION_TITLES):
        _section_heading(document, title)
        if index == 0:
            _basic_info(document, data)
        elif index == 1:
            _device_incidents(document, data)
        elif index == 2:
            _irf_section(document, data)
        elif index == 3:
            _cpu_memory(document, data)
        elif index == 4:
            _interface_incidents(document, data)
        elif index == 5:
            _utilization(document, data)
        elif index == 6:
            _counters_coverage(document, data)
        else:
            _open_questions(document)
    return document


def _validate_document(path: Path) -> None:
    """§5 successful close/validation: reopen and pin the 8 sections."""

    document = Document(str(path))
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style is not None and paragraph.style.name == "Heading 1"
    ]
    if headings != list(SECTION_TITLES):
        raise DocxRenderError(
            f"rendered DOCX does not contain the fixed 8 sections: {headings}"
        )


def render_report_candidate(data: WeeklyReportData, output_dir: Path) -> RenderedReport:
    """Render and validate the weekly DOCX as a candidate, §5 (§4.4).

    The temporary file lives in `output_dir` itself, so the rename to the
    deterministic candidate path is an atomic same-filesystem move. The
    current report of the week is NOT touched here — `install_report`
    does that as its own explicit, atomic step, so a candidate that is
    never installed can never destroy the previous report.
    """

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / report_file_name(data.period)
    candidate = output_dir / candidate_file_name(data.period)

    handle, temp_name = tempfile.mkstemp(
        prefix=".tmp-report-", suffix=".docx", dir=output_dir
    )
    os.close(handle)
    temp_path = Path(temp_name)
    try:
        document = _build_document(data)
        document.save(str(temp_path))
        _validate_document(temp_path)
        os.replace(temp_path, candidate)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    logger.info(
        "weekly report candidate rendered: %s (week %s)", candidate.name, data.period.week_code
    )
    return RenderedReport(path=target, week_code=data.period.week_code, candidate=candidate)


def install_report(rendered: RenderedReport) -> None:
    """Atomically make the candidate the current report of the week (§5).

    Same-directory rename: readers can only ever see the complete previous
    or the complete new document, and the previous bytes stay intact until
    this replace succeeds.
    """

    os.replace(rendered.candidate, rendered.path)


def render_report_docx(data: WeeklyReportData, output_dir: Path) -> RenderedReport:
    """Render, validate and atomically install the weekly DOCX (§5/§4.4).

    One-shot composition of `render_report_candidate` + `install_report`
    for callers that own the whole succeed-or-fail flow in-process.
    """

    rendered = render_report_candidate(data, output_dir)
    install_report(rendered)
    return rendered
