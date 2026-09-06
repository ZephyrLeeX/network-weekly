"""Unit tests for the DOCX renderer (W03-T008, §5/§6).

These tests render a fully-populated and a mostly-empty `WeeklyReportData`
into a temporary directory and validate the produced document by reading
it back with python-docx — no database involved (the statistics layer has
its own tests; the renderer only shapes values into sections/tables).
"""

from datetime import UTC, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from docx import Document

from backend.reporting import docx as docx_module

# The renderer consumes WeeklyReportData; constructing one through the real
# service needs a database. For renderer tests we build the structure with
# the exact production dataclasses via a tiny helper here.
from backend.reporting.counters_coverage import (  # noqa: E402
    CoverageSummary,
    DeviceCoverage,
    InterfaceCounterDelta,
)
from backend.reporting.docx import (
    MISSING,
    SECTION_TITLES,
    DocxRenderError,
    _fmt_duration,
    _validate_document,
    candidate_file_name,
    candidate_week_code,
    install_report,
    is_unbound_candidate_name,
    parse_candidate_name,
    render_report_candidate,
    render_report_docx,
)
from backend.reporting.incidents import DeviceIncidentSummary, IncidentEpisode  # noqa: E402
from backend.reporting.irf_summary import (  # noqa: E402
    IrfDeviceWeeklySummary,
    IrfMemberWeeklySummary,
    IrfMissingWindow,
)
from backend.reporting.period import (
    period_for_iso_week,  # noqa: E402
    report_file_name,
)
from backend.reporting.resources import InterfaceTopEntry  # noqa: E402
from backend.reporting.service import (  # noqa: E402
    STATUS_ABNORMAL,
    STATUS_ATTENTION,
    STATUS_NORMAL,
    DeviceResourceReport,
    WeeklyReportData,
)

PERIOD = period_for_iso_week(2026, 36)
SH_TZ = ZoneInfo("Asia/Shanghai")
GENERATED_AT = datetime(2026, 9, 7, 0, 12, 0, tzinfo=UTC)


def _report(
    *,
    status: str = STATUS_NORMAL,
    summary_text: str = "2026-W36 周报总体状态：正常。",
    coverage: CoverageSummary | None = None,
    device_resources: tuple = (),
    interface_top10: tuple = (),
    device_incidents: tuple = (),
    interface_incidents: tuple = (),
    irf_summaries: tuple = (),
    counter_top10: tuple = (),
    high_utilization: tuple = (),
) -> WeeklyReportData:
    return WeeklyReportData(
        period=PERIOD,
        generated_at=GENERATED_AT,
        timezone=SH_TZ,
        overall_status=status,
        summary_text=summary_text,
        coverage=coverage
        or CoverageSummary(
            devices=(
                DeviceCoverage(
                    device_id=1,
                    device_name="core-1",
                    expected=2016,
                    success=2010,
                    partial=4,
                    failed=2,
                ),
            )
        ),
        device_resources=tuple(device_resources),
        interface_top10=tuple(interface_top10),
        device_incidents=tuple(device_incidents),
        interface_incidents=tuple(interface_incidents),
        irf_summaries=tuple(irf_summaries),
        counter_top10=tuple(counter_top10),
        high_utilization=tuple(high_utilization),
    )


def test_renderer_produces_fixed_eight_sections(tmp_path: Path) -> None:
    path = render_report_docx(_report(), tmp_path, job_id=7).path
    assert path.name == report_file_name(PERIOD)
    assert path.exists()

    document = Document(str(path))
    headings = [
        p.text
        for p in document.paragraphs
        if p.style is not None and p.style.name == "Heading 1"
    ]
    assert headings == list(SECTION_TITLES)
    assert len(SECTION_TITLES) == 8


def test_basic_info_section_contains_required_fields(tmp_path: Path) -> None:
    path = render_report_docx(_report(), tmp_path, job_id=7).path
    document = Document(str(path))
    tables = document.tables
    first = tables[0]
    cells = {(row.cells[0].text, row.cells[1].text) for row in first.rows}
    assert ("报告周编号", "2026-W36") in cells
    assert ("统计开始时间", "2026-08-31 00:00") in cells
    assert ("统计结束时间", "2026-09-07 00:00") in cells
    assert ("报告生成时间", "2026-09-07 08:12") in cells  # UTC 00:12 -> +08
    assert ("当前总体状态", STATUS_NORMAL) in cells
    coverage_row = next(text for label, text in cells if label == "Monitoring Coverage 摘要")
    assert "99.90%" in coverage_row  # (2010+4)/2016 = 99.90%
    assert "PARTIAL 4" in coverage_row
    # The summary paragraph is present as text.
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "2026-W36 周报总体状态：正常。" in all_text


def test_missing_values_render_as_data_missing(tmp_path: Path) -> None:
    report = _report(
        device_resources=(
            DeviceResourceReport(
                device_id=1,
                device_name="silent-core",
                cpu=None,
                memory=None,
                cpu_sustained_high=(),
                memory_sustained_high=(),
            ),
        ),
    )
    path = render_report_docx(report, tmp_path, job_id=7).path
    document = Document(str(path))
    # Section 4 table (index 1 overall: table 0 = basic info): CPU/memory row.
    cpu_table = document.tables[1]
    silent_row = next(row for row in cpu_table.rows if row.cells[0].text == "silent-core")
    stat_cells = [cell.text for cell in silent_row.cells[1:7]]  # avg/max/P95 x2
    assert all(text == MISSING for text in stat_cells)  # §6.2, never 0
    # Interval counts are genuine zeros, not missing.
    assert [cell.text for cell in silent_row.cells[7:]] == ["0", "0"]


def test_incident_tables_and_counter_rows_render(tmp_path: Path) -> None:
    episode = IncidentEpisode(
        started_at=datetime(2026, 9, 1, 8, 0, tzinfo=UTC),
        recovered_at=datetime(2026, 9, 1, 10, 30, tzinfo=UTC),
        started_before_period=False,
        recovered_after_period_end=False,
    )
    device_summary = DeviceIncidentSummary(
        device_id=1, device_name="core-1", episodes=(episode,)
    )
    delta = InterfaceCounterDelta(
        interface_id=1,
        device_id=1,
        device_name="core-1",
        display_name="XGE1/0/1",
        normalized_name="xge1/0/1",
        description=None,
        sample_count=2000,
        crc_delta=550,
        in_errors_delta=None,
        out_errors_delta=0,
        in_discards_delta=None,
        out_discards_delta=None,
    )
    report = _report(
        status=STATUS_ATTENTION,
        summary_text="设备掉线 1 次",
        device_incidents=(device_summary,),
        counter_top10=(delta,),
    )
    path = render_report_docx(report, tmp_path, job_id=7).path
    document = Document(str(path))

    # Section 2: device incident table.
    incident_table = next(
        table
        for table in document.tables
        if table.rows[0].cells[0].text == "设备"
        and table.rows[0].cells[-1].text == "状态"
    )
    row = incident_table.rows[1]
    assert row.cells[0].text == "core-1"
    assert row.cells[3].text == "2小时30分钟"
    assert row.cells[4].text == "已恢复"

    # Section 7: counter delta table with missing columns explicit.
    counter_table = next(
        table
        for table in document.tables
        if len(table.rows[0].cells) > 2 and table.rows[0].cells[2].text == "CRC 增量"
    )
    counter_row = counter_table.rows[1]
    assert counter_row.cells[2].text == "550"  # CRC present
    assert counter_row.cells[3].text == MISSING  # in_errors never reported
    assert counter_row.cells[4].text == "0"  # genuine zero delta


def test_irf_and_open_questions_sections(tmp_path: Path) -> None:
    fabric = IrfDeviceWeeklySummary(
        device_id=2,
        device_name="irf-core",
        expected_member_count=2,
        observed_member_count=1,
        latest_observation_at=datetime(2026, 9, 6, 23, 50, tzinfo=UTC),
        members=(
            IrfMemberWeeklySummary(
                member_id=1,
                observations_count=600,
                missing_windows=(),
                latest_observed=True,
                latest_role="Master",
            ),
            IrfMemberWeeklySummary(
                member_id=2,
                observations_count=600,
                missing_windows=(
                    IrfMissingWindow(
                        started_at=datetime(2026, 9, 3, 8, 0, tzinfo=UTC),
                        reappeared_at=None,
                    ),
                ),
                latest_observed=False,
                latest_role=None,
            ),
        ),
        role_changes=(),
    )
    report = _report(status=STATUS_ABNORMAL, irf_summaries=(fabric,))
    path = render_report_docx(report, tmp_path, job_id=7).path
    document = Document(str(path))
    irf_table = next(
        table for table in document.tables if table.rows[0].cells[0].text == "成员编号"
    )
    member2 = irf_table.rows[2]
    assert member2.cells[1].text == "缺失"
    assert "2026-09-03 16:00" in member2.cells[2].text  # +08 rendering
    assert "期末仍未恢复" in member2.cells[2].text

    # Section 8 exists with an empty editable table area.
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "手工填写" in all_text
    last_table = document.tables[-1]
    assert len(last_table.rows) == 1 and len(last_table.rows[0].cells) == 1
    assert last_table.rows[0].cells[0].text.strip() == ""


def test_zero_device_coverage_renders_missing_not_healthy(tmp_path: Path) -> None:
    """Audit: an empty deployment shows 未配置设备/数据缺失, never healthy."""

    report = _report(
        status=STATUS_ATTENTION,
        summary_text="2026-W36 周报总体状态：关注。未配置设备/数据缺失；数据完整性不足。",
        coverage=CoverageSummary(devices=()),
    )
    path = render_report_docx(report, tmp_path, job_id=7).path
    document = Document(str(path))
    cells = {(row.cells[0].text, row.cells[1].text) for row in document.tables[0].rows}
    assert ("当前总体状态", STATUS_ATTENTION) in cells  # never 正常 over an empty DB
    coverage_row = next(text for label, text in cells if label == "Monitoring Coverage 摘要")
    assert "未配置设备/数据缺失" in coverage_row
    assert "数据完整性不足" in coverage_row
    assert "总体" not in coverage_row  # no percentage over zero devices
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "数据完整性满足要求" not in all_text
    assert "未配置设备/数据缺失" in all_text


def test_irf_data_missing_fabric_renders_missing_not_loss(tmp_path: Path) -> None:
    """Audit: no successful observation reads as 数据缺失, not member loss."""

    fabric = IrfDeviceWeeklySummary(
        device_id=3,
        device_name="irf-core",
        expected_member_count=2,
        observed_member_count=None,
        latest_observation_at=None,
        members=(
            IrfMemberWeeklySummary(
                member_id=1,
                observations_count=0,
                missing_windows=(),
                latest_observed=None,
                latest_role=None,
            ),
            IrfMemberWeeklySummary(
                member_id=2,
                observations_count=0,
                missing_windows=(),
                latest_observed=None,
                latest_role=None,
            ),
        ),
        role_changes=(),
    )
    report = _report(
        status=STATUS_ATTENTION,
        summary_text="2026-W36 周报总体状态：关注。IRF 成员状态数据缺失（irf-core）。",
        irf_summaries=(fabric,),
    )
    path = render_report_docx(report, tmp_path, job_id=7).path
    document = Document(str(path))
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "本周无成功观察，成员状态数据缺失" in all_text  # IRF section header
    assert "IRF 成员状态数据缺失（irf-core）" in all_text  # summary paragraph
    irf_table = next(
        table for table in document.tables if table.rows[0].cells[0].text == "成员编号"
    )
    assert [row.cells[1].text for row in irf_table.rows[1:]] == [MISSING, MISSING]


def test_top10_table_renders_rank_and_utilization(tmp_path: Path) -> None:
    entry = InterfaceTopEntry(
        interface_id=1,
        device_id=1,
        device_name="core-1",
        display_name="XGE1/0/1",
        normalized_name="xge1/0/1",
        description="uplink",
        sample_count=2000,
        average=55.0,
        maximum=88.4,
        p95=80.25,
    )
    report = _report(interface_top10=(entry,))
    path = render_report_docx(report, tmp_path, job_id=7).path
    document = Document(str(path))
    top_table = next(
        table for table in document.tables if table.rows[0].cells[0].text == "排名"
    )
    row = top_table.rows[1]
    assert row.cells[0].text == "1"
    assert row.cells[1].text == "core-1"
    assert row.cells[2].text == "XGE1/0/1"
    assert row.cells[5].text == "80.25%"
    assert row.cells[6].text == "88.4%"


def test_atomic_replace_leaves_exactly_one_file_per_week(tmp_path: Path) -> None:
    first = render_report_docx(_report(), tmp_path, job_id=7)
    assert [p.name for p in tmp_path.iterdir()] == [first.path.name]
    # Regenerate over the same week: one current file, no temp leftovers.
    second = render_report_docx(_report(), tmp_path, job_id=7)
    assert second.path == first.path
    assert [p.name for p in tmp_path.iterdir()] == [first.path.name]


def test_candidate_then_install_is_two_phase(tmp_path: Path) -> None:
    """§4.4: the candidate is a complete validated DOCX bound to its
    generating job; the current report appears only when the install
    explicitly switches it."""

    rendered = render_report_candidate(_report(), tmp_path, job_id=42)
    assert rendered.path.name == report_file_name(PERIOD)
    assert rendered.job_id == 42
    assert rendered.candidate.name == candidate_file_name(PERIOD, 42)
    # The candidate holds the complete document; the current report of the
    # week was not touched by the render.
    assert rendered.candidate.exists() and not rendered.path.exists()
    assert [p.name for p in tmp_path.iterdir()] == [rendered.candidate.name]
    document = Document(str(rendered.candidate))
    headings = [
        p.text for p in document.paragraphs if p.style is not None and p.style.name == "Heading 1"
    ]
    assert headings == list(SECTION_TITLES)

    install_report(rendered)
    assert rendered.path.exists() and not rendered.candidate.exists()
    assert [p.name for p in tmp_path.iterdir()] == [rendered.path.name]


def test_candidate_week_code_round_trip() -> None:
    name = candidate_file_name(PERIOD, 12)
    assert candidate_week_code(name) == "2026-W36"
    assert parse_candidate_name(name) == ("2026-W36", 12)
    # Only job-bound candidate files parse; current reports and foreign
    # names do not — and a legacy unbound candidate never parses either.
    assert candidate_week_code(report_file_name(PERIOD)) is None
    assert candidate_week_code("network-weekly-report-2026-W36.docx.candidate") is None
    assert candidate_week_code("network-weekly-report-2026-W366.docx.candidate.1") is None
    assert candidate_week_code("network-weekly-report-2026-W36.docx.candidate.1x") is None
    assert candidate_week_code("someone-elses-file.candidate") is None
    assert is_unbound_candidate_name("network-weekly-report-2026-W36.docx.candidate")
    assert not is_unbound_candidate_name(name)
    assert not is_unbound_candidate_name("someone-elses-file.candidate")


def test_invalid_document_fails_validation(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """A corrupted save never reaches the target (temp file -> validate -> replace)."""

    def broken_build(_data: object) -> object:
        raise RuntimeError("boom")

    monkeypatch.setattr(docx_module, "_build_document", broken_build)
    with pytest.raises(RuntimeError, match="boom"):
        render_report_docx(_report(), tmp_path, job_id=7)
    assert list(tmp_path.iterdir()) == []  # no partial artifacts


def test_validate_document_rejects_missing_sections(tmp_path: Path) -> None:
    document = Document()
    document.add_heading(SECTION_TITLES[0], level=1)  # only 1 of 8
    path = tmp_path / "broken.docx"
    document.save(str(path))
    with pytest.raises(DocxRenderError, match="fixed 8 sections"):
        _validate_document(path)


def test_duration_formatting() -> None:
    assert _fmt_duration(None) == "未恢复"
    assert _fmt_duration(900.0) == "15分钟"
    assert _fmt_duration(9000.0) == "2小时30分钟"
    assert _fmt_duration(90000.0) == "1天1小时"
    assert _fmt_duration(59.0) == "0分钟"  # below display granularity
