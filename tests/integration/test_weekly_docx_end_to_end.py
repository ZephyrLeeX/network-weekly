"""End-to-end integration test: persisted data -> statistics -> DOCX
(W03-T008 acceptance) over migrated PostgreSQL.
"""

from collections.abc import Iterator
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import Engine
from sqlalchemy.orm import Session

from backend.db.models import (
    Device,
    DeviceMember,
    DevicePollRun,
    DeviceReachabilityIncident,
)
from backend.reporting.docx import MISSING, SECTION_TITLES, render_report_docx
from backend.reporting.period import period_for_iso_week
from backend.reporting.service import (
    STATUS_ATTENTION,
    build_weekly_report_data,
)

pytestmark = pytest.mark.integration

PERIOD = period_for_iso_week(2026, 36)


@pytest.fixture(autouse=True)
def _clean(db_engine: Engine) -> Iterator[None]:
    with Session(db_engine) as session:
        for model in (DevicePollRun, DeviceReachabilityIncident, Device):
            session.query(model).delete()
        session.commit()
        yield


def test_report_chain_renders_valid_docx_with_missing_data(
    db_engine: Engine, tmp_path: Path
) -> None:
    with Session(db_engine) as session:
        # One device with a couple of cycles, one silent device, one open
        # incident: report must still generate with explicit 数据缺失.
        talking = Device(
            name="core-talk",
            management_ip="192.0.2.21",
            model_family="s10500x",
            credential_profile="default",
        )
        silent = Device(
            name="core-silent",
            management_ip="192.0.2.22",
            model_family="s10500x",
            credential_profile="default",
        )
        session.add_all([talking, silent])
        session.flush()
        for i in range(3):
            session.add(
                DevicePollRun(
                    device_id=talking.id,
                    cycle_started_at=PERIOD.start + timedelta(minutes=5 * i),
                    status="SUCCESS",
                )
            )
        session.add(
            DeviceReachabilityIncident(
                device_id=silent.id,
                started_at=datetime(2026, 9, 3, tzinfo=UTC),
                recovered_at=None,
            )
        )
        session.commit()

        generated = datetime(2026, 9, 7, 0, 12, tzinfo=UTC)
        data = build_weekly_report_data(session, PERIOD, generated_at=generated)

    rendered = render_report_docx(data, tmp_path)
    assert rendered.path.name == "network-weekly-report-2026-W36.docx"

    document = Document(str(rendered.path))
    headings = [
        p.text
        for p in document.paragraphs
        if p.style is not None and p.style.name == "Heading 1"
    ]
    assert headings == list(SECTION_TITLES)

    body_text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                body_text += "\n" + cell.text
    assert MISSING in body_text  # silent device CPU/memory explicitly missing
    assert "仍处于 Down" in body_text  # ongoing incident visible
    assert "2026-W36" in body_text
    assert "数据完整性不足" in body_text  # far below 95% coverage

    # 本周处理问题 blank editable area: last table is an empty 1x1 grid cell.
    last_table = document.tables[-1]
    assert len(last_table.rows) == 1 and len(last_table.rows[0].cells) == 1
    assert last_table.rows[0].cells[0].text.strip() == ""

    # No secret-bearing material ever appears in the document (§22.2).
    for marker in ("community", "password", "secret", "snmp", "credential"):
        assert marker not in body_text.lower()


def test_zero_device_database_renders_missing_configuration(
    db_engine: Engine, tmp_path: Path
) -> None:
    """Audit: a truly empty database (0 devices) is 数据缺失, never healthy."""

    with Session(db_engine) as session:
        assert session.query(Device).count() == 0  # fixture left no device behind
        generated = datetime(2026, 9, 7, 0, 12, tzinfo=UTC)
        data = build_weekly_report_data(session, PERIOD, generated_at=generated)

    assert data.overall_status == STATUS_ATTENTION  # never 正常 over an empty DB
    rendered = render_report_docx(data, tmp_path)
    document = Document(str(rendered.path))
    assert [p.text for p in document.paragraphs if p.style and p.style.name == "Heading 1"] == (
        list(SECTION_TITLES)
    )

    body_text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                body_text += "\n" + cell.text
    assert "未配置设备/数据缺失" in body_text
    assert "数据完整性满足要求" not in body_text
    assert "数据完整性不足" in body_text

    basic_info = {(row.cells[0].text, row.cells[1].text) for row in document.tables[0].rows}
    assert ("当前总体状态", STATUS_ATTENTION) in basic_info
    coverage_row = next(text for label, text in basic_info if label == "Monitoring Coverage 摘要")
    assert "未配置设备/数据缺失" in coverage_row

    # Section 7 coverage table: overall row is explicit 数据缺失, not 0%.
    coverage_table = document.tables[-2]
    overall = coverage_table.rows[-1]
    assert overall.cells[0].text == "总体"
    assert overall.cells[5].text == MISSING
    assert overall.cells[6].text == "数据完整性不足"


def test_irf_data_missing_week_renders_missing_not_loss(
    db_engine: Engine, tmp_path: Path
) -> None:
    """Audit: a fabric with no successful observation states 数据缺失."""

    with Session(db_engine) as session:
        fabric = Device(
            name="core-irf",
            management_ip="192.0.2.30",
            model_family="s10500x",
            expected_irf_member_count=2,
            credential_profile="default",
        )
        session.add(fabric)
        session.flush()
        session.add_all(
            [
                DeviceMember(device_id=fabric.id, member_id=1),
                DeviceMember(device_id=fabric.id, member_id=2),
            ]
        )
        session.add(
            DevicePollRun(device_id=fabric.id, cycle_started_at=PERIOD.start, status="SUCCESS")
        )
        session.commit()

        generated = datetime(2026, 9, 7, 0, 12, tzinfo=UTC)
        data = build_weekly_report_data(session, PERIOD, generated_at=generated)

    assert data.irf_summaries[0].data_missing
    assert data.overall_status == STATUS_ATTENTION  # §19 + Coverage, not 异常
    rendered = render_report_docx(data, tmp_path)
    document = Document(str(rendered.path))

    body_text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                body_text += "\n" + cell.text
    assert "IRF 成员状态数据缺失（core-irf）" in body_text  # overall summary
    assert "IRF 成员无缺失" not in body_text
    assert "本周无成功观察，成员状态数据缺失" in body_text  # IRF section
    assert "总体状态：异常" not in body_text  # missing data is not an 异常 verdict
